"""
Turns a report (see reports.py) into an Excel file or a PDF.
The PDF look is in backend/templates/documents/report.html (SAFE TO EDIT there).
"""
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.template.loader import render_to_string
from django.utils import timezone


def _text(value, kind):
    """A cell as text for the PDF: 1,25,050.50 style money, 06-10-2026 dates, Oct 2026 months."""
    if value is None or value == "":
        return ""
    if kind == "money" and isinstance(value, Decimal):
        return indian_money(value)
    if kind == "date" and isinstance(value, date):
        return value.strftime("%d-%m-%Y")
    if kind == "month" and isinstance(value, date):
        return value.strftime("%b %Y")
    return str(value)


def indian_money(value: Decimal) -> str:
    """12345678.5 -> 1,23,45,678.50 (Indian grouping)."""
    negative = value < 0
    whole, paise = f"{abs(value):.2f}".split(".")
    head, tail = whole[:-3], whole[-3:]
    groups = []
    while len(head) > 2:
        groups.insert(0, head[-2:])
        head = head[:-2]
    if head:
        groups.insert(0, head)
    text = ",".join(groups + [tail]) if groups else tail
    return f"{'-' if negative else ''}{text}.{paise}"


def to_excel(report: dict, *, title: str, subtitle: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    columns = report["columns"]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([subtitle])
    ws.append([])
    ws.append([c["label"] for c in columns])
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEF3F0")
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    def put(row, bold=False):
        values = []
        for c in columns:
            v = row.get(c["key"])
            values.append(float(v) if isinstance(v, Decimal) else v)
        ws.append(values)
        for c, cell in zip(columns, ws[ws.max_row]):
            if c["type"] == "money":
                cell.number_format = "#,##0.00"
            elif c["type"] == "date":
                cell.number_format = "DD-MM-YYYY"
            elif c["type"] == "month":
                cell.number_format = "MMM YYYY"
            if bold:
                cell.font = Font(bold=True)

    for row in report["rows"]:
        put(row)
    if report.get("totals"):
        put(report["totals"], bold=True)
    for i, c in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 26 if c["type"] == "text" else 15
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(report: dict, *, title: str, subtitle: str, organization, notes: str = "") -> bytes:
    from weasyprint import HTML

    columns = report["columns"]
    landscape = len(columns) > 6
    rows = [[{"text": _text(r.get(c["key"]), c["type"]), "number": c["type"] in ("money", "int")} for c in columns]
            for r in report["rows"]]
    totals = report.get("totals")
    html = render_to_string("documents/report.html", {
        "title": title, "subtitle": subtitle, "organization": organization, "notes": notes,
        "columns": [{"label": c["label"], "number": c["type"] in ("money", "int")} for c in columns],
        "rows": rows,
        "totals": [{"text": _text(totals.get(c["key"]), c["type"]), "number": c["type"] in ("money", "int")}
                   for c in columns] if totals else None,
        "printed_at": timezone.localtime().strftime("%d-%m-%Y %H:%M"),
        "page_size": "A4 landscape" if landscape else "A4",
    })
    return HTML(string=html).write_pdf()
