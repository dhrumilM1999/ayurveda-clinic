"""
Medicine list rules: saving with version history, searching by any name, branch price, Excel/CSV import.
Other modules (prescriptions, pharmacy, billing) should use these functions.
"""
import csv
import io
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.db.models import Q
from rest_framework.exceptions import ValidationError

from apps.common.models import MasterValue

from .models import BranchMedicine, Medicine, MedicineVersion

# Fields copied into the version history (and offered in the import file)
TEXT_FIELDS = [
    "kind", "name", "name_gu", "name_hi", "synonyms", "generic_name", "composition", "reference", "manufacturer",
    "ayush_licence_no", "hsn_code", "pack_size", "default_dose", "default_frequency", "safety_notes", "barcode",
    "strength", "sku", "notes",
]
MASTER_FIELDS = {  # field -> dropdown list
    "dosage_form": "dosage_form", "dose_unit": "dose_unit",
    "default_timing": "medicine_timing", "default_anupana": "anupana",
    "category": "product_category", "pack_type": "pack_type",
}
FLAG_FIELDS = ["schedule_e1", "contains_metals", "pregnancy_caution", "child_caution", "allow_loose"]
NUMBER_FIELDS = ["gst_rate", "mrp", "selling_price", "units_per_pack"]


# --- Search ---------------------------------------------------------------------
def search_filter(text: str) -> Q:
    """Find by name in any language, synonyms, composition or manufacturer."""
    q = Q()
    for word in text.split():
        q &= (
            Q(name__icontains=word) | Q(name_gu__icontains=word) | Q(name_hi__icontains=word)
            | Q(synonyms__icontains=word) | Q(composition__icontains=word) | Q(manufacturer__icontains=word)
            | Q(generic_name__icontains=word) | Q(barcode=word) | Q(sku__iexact=word)
        )
    return q


# --- Version history -------------------------------------------------------------
def snapshot(medicine: Medicine) -> dict:
    data = {f: getattr(medicine, f) for f in TEXT_FIELDS + FLAG_FIELDS}
    for f in NUMBER_FIELDS:
        value = getattr(medicine, f)
        data[f] = str(value) if value is not None else None
    for f in MASTER_FIELDS:
        value = getattr(medicine, f)
        data[f] = value.label if value else ""
    data["classical_equivalent"] = medicine.classical_equivalent.name if medicine.classical_equivalent_id else ""
    data["is_active"] = medicine.is_active
    return data


def record_version(medicine: Medicine, user=None):
    MedicineVersion.objects.create(medicine=medicine, version=medicine.version, data=snapshot(medicine),
                                   created_by=user)


def same_details(old: dict, new: dict) -> bool:
    """True if nothing that matters changed. Fields added later (missing in old versions) count as empty."""
    return all(old.get(k, "") == v for k, v in new.items())


@transaction.atomic
def save_medicine(medicine: Medicine, user, is_new: bool) -> Medicine:
    """Save; if anything that matters changed, raise the version and keep a copy of the old details."""
    if is_new:
        medicine.version = 1
        medicine.created_by = medicine.updated_by = user
        medicine.save()
        record_version(medicine, user)
        return medicine
    previous = MedicineVersion.objects.filter(medicine=medicine).order_by("-version").first()
    medicine.updated_by = user
    if previous is None or not same_details(previous.data, snapshot(medicine)):
        medicine.version += 1
        medicine.save()
        record_version(medicine, user)
    else:
        medicine.save()
    return medicine


# --- Branch price / on-off -----------------------------------------------------------
def set_branch_settings(medicine: Medicine, branch, user, price=None, is_active=None) -> BranchMedicine:
    row = BranchMedicine.objects.filter(branch=branch, medicine=medicine).first()
    if row is None:
        row = BranchMedicine(organization_id=medicine.organization_id, branch=branch, medicine=medicine,
                             created_by=user)
    if price is not None:
        row.price = None if price == "" else price
    if is_active is not None:
        row.is_active = is_active
    row.updated_by = user
    row.save()
    return row


def branch_price(medicine: Medicine, branch):
    row = BranchMedicine.objects.filter(branch=branch, medicine=medicine).first()
    return row.price if row and row.price is not None else medicine.mrp


# --- Import (Excel or CSV) ------------------------------------------------------------
IMPORT_COLUMNS = TEXT_FIELDS + list(MASTER_FIELDS) + NUMBER_FIELDS + FLAG_FIELDS + ["classical_equivalent"]
YES = {"yes", "y", "true", "1", "હા", "हाँ", "हां"}


def import_template_csv() -> str:
    """An example file with the right column names."""
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(IMPORT_COLUMNS)
    example = {
        "kind": "classical", "name": "Triphala Churna", "name_gu": "ત્રિફળા ચૂર્ણ", "name_hi": "त्रिफला चूर्ण",
        "synonyms": "Triphala, Three fruits", "composition": "Haritaki, Bibhitaki, Amalaki",
        "reference": "AFI Part I", "hsn_code": "3004", "pack_size": "100 g", "default_dose": "3",
        "default_frequency": "0-0-1", "dosage_form": "churna", "dose_unit": "g", "default_timing": "bedtime",
        "default_anupana": "warm_water", "gst_rate": "12", "mrp": "90", "schedule_e1": "no",
        "contains_metals": "no", "pregnancy_caution": "no", "child_caution": "no",
    }
    writer.writerow([example.get(c, "") for c in IMPORT_COLUMNS])
    return out.getvalue()


def read_rows(uploaded) -> list[dict]:
    """Rows of an uploaded .csv or .xlsx file as {column: text}."""
    name = (uploaded.name or "").lower()
    if uploaded.size > 5 * 1024 * 1024:
        raise ValidationError({"file": "File is too large (maximum 5 MB)."})
    if name.endswith(".csv"):
        text = uploaded.read().decode("utf-8-sig", errors="replace")
        return [dict(row) for row in csv.DictReader(io.StringIO(text))]
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            sheet = load_workbook(uploaded, read_only=True, data_only=True).active
        except Exception:
            raise ValidationError({"file": "This Excel file could not be read."})
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h or "").strip() for h in rows[0]]
        return [
            {header[i]: ("" if v is None else str(v)) for i, v in enumerate(r) if i < len(header)}
            for r in rows[1:] if any(v not in (None, "") for v in r)
        ]
    raise ValidationError({"file": "Please upload a .xlsx (Excel) or .csv file."})


def _master_lookup(organization, category):
    lookup = {}
    for m in MasterValue.objects.filter(organization=organization, category=category):
        for text in (m.code, m.label, m.label_gu, m.label_hi):
            if text:
                lookup[text.strip().lower()] = m
    return lookup


def import_medicines(organization, rows: list[dict], user, dry_run: bool) -> dict:
    """
    Add or update medicines from file rows (matched by kind + name). With dry_run nothing is saved:
    the result shows what WOULD happen, so the user can check first.
    """
    masters = {field: _master_lookup(organization, cat) for field, cat in MASTER_FIELDS.items()}
    existing = {
        (m.kind, m.name.strip().lower()): m for m in Medicine.objects.filter(organization=organization)
    }
    result = {"created": 0, "updated": 0, "unchanged": 0, "errors": [], "rows": []}
    pending_equivalents = []

    with transaction.atomic():
        for number, raw in enumerate(rows, start=2):  # row 1 is the header
            row = {str(k or "").strip().lower(): str(v or "").strip() for k, v in raw.items()}
            name = row.get("name", "")
            if not name:
                result["errors"].append({"row": number, "message": "Name is missing."})
                continue
            kind_text = row.get("kind", "classical").lower()
            kind = "proprietary" if kind_text.startswith(("p", "brand")) else "classical"
            try:
                medicine = existing.get((kind, name.lower()))
                is_new = medicine is None
                if is_new:
                    medicine = Medicine(organization=organization, kind=kind, name=name[:200])
                for field in TEXT_FIELDS:
                    if field in row and field not in ("kind", "name"):
                        setattr(medicine, field, row[field][:300] if field != "composition" else row[field])
                for field in MASTER_FIELDS:
                    if row.get(field):
                        match = masters[field].get(row[field].lower())
                        if match is None:
                            raise ValueError(f"'{row[field]}' is not in the {field.replace('_', ' ')} list.")
                        setattr(medicine, field, match)
                for field in NUMBER_FIELDS:
                    if row.get(field):
                        try:
                            setattr(medicine, field, Decimal(row[field]))
                        except InvalidOperation:
                            raise ValueError(f"{field} must be a number.")
                for field in FLAG_FIELDS:
                    if field in row:
                        setattr(medicine, field, row[field].lower() in YES)
                before = None if is_new else MedicineVersion.objects.filter(medicine=medicine).order_by("-version").first()
                changed = is_new or before is None or not same_details(before.data, snapshot(medicine))
                if not dry_run and changed:
                    save_medicine(medicine, user, is_new)
                    existing[(kind, name.lower())] = medicine
                if row.get("classical_equivalent") and kind == "proprietary":
                    pending_equivalents.append((medicine, row["classical_equivalent"], number))
                status = "created" if is_new else ("updated" if changed else "unchanged")
                result[status] += 1
                result["rows"].append({"row": number, "name": name, "kind": kind, "status": status})
            except ValueError as error:
                result["errors"].append({"row": number, "message": f"{name}: {error}"})

        for medicine, classical_name, number in pending_equivalents:
            match = existing.get(("classical", classical_name.lower()))
            if match is None:
                result["errors"].append({"row": number, "message": f"Classical medicine '{classical_name}' not found."})
            elif not dry_run and medicine.classical_equivalent_id != match.id:
                medicine.classical_equivalent = match
                save_medicine(medicine, user, is_new=False)
        if dry_run:
            transaction.set_rollback(True)
    return result


# --- Sample medicines for a new clinic -----------------------------------------------------
def add_sample_medicines(organization) -> int:
    from .sample_catalog import EQUIVALENTS, MEDICINES

    masters = {field: _master_lookup(organization, cat) for field, cat in MASTER_FIELDS.items()}
    added = {}
    for (kind, name, gu, hi, synonyms, form, composition, reference, dose, unit, freq, timing, anupana,
         mrp, pack, flags) in MEDICINES:
        if Medicine.objects.filter(organization=organization, kind=kind, name=name).exists():
            continue
        flag_set = {f.strip() for f in flags.split(",") if f.strip()}
        medicine = Medicine(
            organization=organization, kind=kind, name=name, name_gu=gu, name_hi=hi, synonyms=synonyms,
            dosage_form=masters["dosage_form"].get(form), composition=composition, reference=reference,
            manufacturer="Sample Pharma (made up)" if kind == "proprietary" else "",
            default_dose=dose, dose_unit=masters["dose_unit"].get(unit), default_frequency=freq,
            default_timing=masters["default_timing"].get(timing), default_anupana=masters["default_anupana"].get(anupana),
            mrp=mrp, pack_size=pack, hsn_code="3004", gst_rate=12,
            schedule_e1="E1" in flag_set, contains_metals="M" in flag_set,
            pregnancy_caution="P" in flag_set, child_caution="C" in flag_set, is_sample=True,
        )
        save_medicine(medicine, None, is_new=True)
        added[name] = medicine
    for brand, classical in EQUIVALENTS.items():
        if brand in added and classical in added:
            added[brand].classical_equivalent = added[classical]
            added[brand].save(update_fields=["classical_equivalent"])
    return len(added)
